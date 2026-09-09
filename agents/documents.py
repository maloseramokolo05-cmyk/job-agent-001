from __future__ import annotations

import hashlib
import io
import re
from datetime import date

from docx import Document
from docx.shared import Pt

from backend.database import connect, now
from backend.storage import ObjectStorage

GENERATOR_VERSION = "3.0.0"
TEMPLATE_VERSION = "ats-single-column-2"


def safe_filename(value: str) -> str:
    value = re.sub(r'[\\/:*?"<>|]+', "-", value)
    value = re.sub(r"\s+", " ", value).strip(" .-")
    return value[:100] or "Job"


def _facts(cv_text: str) -> list[str]:
    return [line.strip() for line in cv_text.splitlines() if line.strip()]


def _terms(job) -> set[str]:
    stop = {
        "and", "the", "with", "for", "from", "that", "this", "you", "your", "our", "are", "will",
        "have", "has", "job", "role", "work", "company", "candidate", "required", "requirements",
    }
    text = f"{job.get('title', '')} {job.get('requirements', '')} {job.get('description', '')}"
    return {
        word.lower()
        for word in re.findall(r"[A-Za-z][A-Za-z0-9+#.-]{2,}", text)
        if word.lower() not in stop
    }


def _ranked_facts(job, cv_text: str) -> list[str]:
    terms = _terms(job)
    facts = _facts(cv_text)
    scored = []
    for index, line in enumerate(facts):
        lowered = line.lower()
        score = sum(1 for term in terms if term in lowered)
        scored.append((score, -index, line))
    return [line for _score, _index, line in sorted(scored, reverse=True)]


def _evidence_summary(job, cv_text: str) -> str:
    ranked = [line for line in _ranked_facts(job, cv_text) if len(line) > 25]
    evidence = ranked[:3]
    if not evidence:
        return "Candidate background and qualifications are presented below exactly as supported by the master CV."
    joined = " ".join(evidence)
    return joined[:650]


def _store(filename: str, content: bytes, content_type: str) -> tuple[str, dict]:
    key = f"generated/{date.today().isoformat()}/{safe_filename(filename)}"
    metadata = ObjectStorage().put(key, content, content_type)
    return f"storage:{metadata['storage_key']}", metadata


def _record(job_id, kind, ref, metadata, master):
    if not job_id:
        return
    digest = hashlib.sha256(master.encode()).hexdigest()
    with connect() as db:
        db.execute(
            "INSERT INTO documents(job_id,document_type,local_path,storage_key,content_type,created_at,master_cv_hash,template_version,generator_version,content_hash) VALUES(?,?,?,?,?,?,?,?,?,?)",
            (
                job_id,
                kind,
                ref,
                metadata["storage_key"],
                metadata["content_type"],
                now(),
                digest,
                TEMPLATE_VERSION,
                GENERATOR_VERSION,
                metadata["sha256"],
            ),
        )


def generate_cv(job, profile, cv_text, job_id=None):
    if not cv_text.strip():
        raise ValueError("A non-empty verified master CV is required; documents are never fabricated")
    name = profile.get("name") or "Candidate"
    title = safe_filename(job.get("title", "Job"))
    filename_docx = f"{name} - {title} - CV.docx"
    filename_pdf = f"{name} - {title} - CV.pdf"
    contacts = " | ".join(
        value for value in [
            profile.get("email"), profile.get("phone"), profile.get("location"),
            profile.get("linkedin"), profile.get("portfolio"),
        ] if value
    )
    ranked = _ranked_facts(job, cv_text)

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    document.add_heading(name, 0)
    document.add_paragraph(contacts)
    document.add_heading("Professional Profile", 1)
    document.add_paragraph(
        f"Application focus: {job.get('title', 'the advertised role')} at {job.get('company', 'the employer')}. "
        f"Relevant evidence from the verified master CV: {_evidence_summary(job, cv_text)}"
    )
    document.add_heading("Relevant Skills, Experience & Qualifications", 1)
    for line in ranked:
        document.add_paragraph(line, style="List Bullet" if len(line) < 180 else None)
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)
    docx_ref, docx_meta = _store(filename_docx, docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas

    pdf_buffer = io.BytesIO()
    canvas = Canvas(pdf_buffer, pagesize=A4)
    width, height = A4
    y = height - 48
    canvas.setFont("Helvetica-Bold", 16)
    canvas.drawString(42, y, name)
    y -= 20
    canvas.setFont("Helvetica", 8.5)
    for chunk in [contacts[index:index + 110] for index in range(0, len(contacts), 110)] or [""]:
        canvas.drawString(42, y, chunk)
        y -= 12
    y -= 4
    canvas.setFont("Helvetica-Bold", 11)
    canvas.drawString(42, y, "Professional Profile")
    y -= 16
    canvas.setFont("Helvetica", 8.8)
    profile_text = (
        f"Application focus: {job.get('title', 'the advertised role')} at {job.get('company', 'the employer')}. "
        f"{_evidence_summary(job, cv_text)}"
    )
    lines = [profile_text] + ranked
    for line in lines:
        for chunk in [line[index:index + 112] for index in range(0, len(line), 112)] or [""]:
            if y < 48:
                canvas.showPage()
                canvas.setFont("Helvetica", 8.8)
                y = height - 48
            canvas.drawString(42, y, chunk)
            y -= 11
        y -= 3
    canvas.save()
    pdf_ref, pdf_meta = _store(filename_pdf, pdf_buffer.getvalue(), "application/pdf")

    _record(job_id, "CV_DOCX", docx_ref, docx_meta, cv_text)
    _record(job_id, "CV_PDF", pdf_ref, pdf_meta, cv_text)
    return docx_ref, pdf_ref


def generate_cover_letter(job, profile, cv_text, job_id=None):
    if not cv_text.strip():
        raise ValueError("A verified master CV is required")
    name = profile.get("name") or "Candidate"
    title = safe_filename(job.get("title", "Job"))
    filename = f"{name} - {title} - Cover Letter.docx"
    document = Document()
    document.add_paragraph(name)
    document.add_paragraph("Dear Hiring Team,")
    document.add_paragraph(
        f"I am applying for the {job.get('title', 'advertised')} opportunity at {job.get('company', 'your organisation')}. "
        "My application is based only on the experience, qualifications and skills documented in my CV."
    )
    document.add_paragraph(f"Relevant background: {_evidence_summary(job, cv_text)}")
    document.add_paragraph(
        "I would welcome the opportunity to discuss how this background can support the requirements of the role."
    )
    document.add_paragraph(f"Kind regards,\n{name}")
    buffer = io.BytesIO()
    document.save(buffer)
    ref, metadata = _store(filename, buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    _record(job_id, "COVER_LETTER", ref, metadata, cv_text)
    return ref


def read_document(ref: str):
    if not ref:
        raise ValueError("Document reference is empty")
    if ref.startswith("storage:"):
        key = ref.split(":", 1)[1]
        stream, content_type = ObjectStorage().get(key)
        return stream.getvalue(), content_type, key.rsplit("/", 1)[-1]
    from backend.config import ROOT

    path = ROOT / ref
    content_type = "application/pdf" if path.suffix.lower() == ".pdf" else "application/octet-stream"
    return path.read_bytes(), content_type, path.name
